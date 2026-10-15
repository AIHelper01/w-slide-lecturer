#!/usr/bin/env python3
"""Validate guide coverage, image paths, and source/explanation emphasis."""
import argparse
import json
import re
from pathlib import Path

def normalize(text):
    return re.sub(r'\s+', ' ', text).strip()

def validate(data, base, docx=None):
    errors = []
    slides = data.get('slides', [])
    if not slides: errors.append('No slides')
    nums = [s.get('number') for s in slides]
    if not all(isinstance(n, int) and not isinstance(n, bool) and n > 0 for n in nums):
        errors.append('Slide numbers must be positive integers')
    elif nums != sorted(set(nums)): errors.append('Slide numbers must be unique and ascending')
    count = data.get('source_slide_count')
    if count is not None and nums != list(range(1, count + 1)):
        errors.append('Full deck coverage does not match source_slide_count')
    expected_bold, expected_regular = [], []
    for slide in slides:
        n = slide.get('number')
        if not slide.get('title'): errors.append(f'Slide {n}: missing title')
        img = slide.get('image')
        if not img or not (base / img).is_file(): errors.append(f'Slide {n}: missing original image')
        modules = slide.get('modules', [])
        expected_ids = slide.get('expected_source_ids')
        if expected_ids is not None:
            actual_ids = [sid for m in modules for sid in m.get('source_ids', [])]
            if len(actual_ids) != len(set(actual_ids)) or set(actual_ids) != set(expected_ids):
                errors.append(f'Slide {n}: missing or duplicate original source units')
        if not modules: errors.append(f'Slide {n}: no modules')
        for mod in modules:
            label = mod.get('label', '?')
            revised = mod.get('revised', [])
            if not isinstance(revised, list) or not all(isinstance(t, str) for t in revised):
                errors.append(f'Slide {n}/{label}: revised must be strings'); continue
            reference = normalize(' '.join(revised))
            if not reference: errors.append(f'Slide {n}/{label}: no revised content')
            original = mod.get('original')
            if original:
                old_words = len(normalize(' '.join(original)).split())
                new_words = len(reference.split())
                if old_words >= 20 and not .9 <= new_words / old_words <= 1.1:
                    errors.append(f'Slide {n}/{label}: length changed over 10%; review for omissions or added source content')
            source_count = 0
            script = mod.get('script', [])
            if not script: errors.append(f'Slide {n}/{label}: no script')
            for para in script:
                if not isinstance(para, list) or not para:
                    errors.append(f'Slide {n}/{label}: empty or invalid script paragraph'); continue
                for seg in para:
                    text, flag = seg.get('text'), seg.get('source')
                    if not isinstance(text, str) or type(flag) is not bool:
                        errors.append(f'Slide {n}/{label}: segment requires text and boolean source'); continue
                    if flag:
                        source_count += 1
                        if not normalize(text) or normalize(text) not in reference:
                            errors.append(f'Slide {n}/{label}: source text not found in revised wording')
                        expected_bold.append(text)
                    else: expected_regular.append(text)
            if not source_count: errors.append(f'Slide {n}/{label}: no bold source segment')
    if docx:
        from docx import Document
        doc = Document(docx)
        heads = [p.text for p in doc.paragraphs if p.style.name == 'Heading 1' and p.text.startswith('Slide ')]
        expected_heads = [f"Slide {s['number']:02d} — {s['title']}" for s in slides]
        if heads != expected_heads: errors.append('DOCX slide headings mismatch')
        if len(doc.inline_shapes) < len(slides): errors.append('DOCX missing original slide images')
        runs = [r for p in doc.paragraphs for r in p.runs]
        for text in expected_bold:
            if not any(r.text == text and r.bold is True for r in runs):
                errors.append('DOCX missing bold source segment: ' + text[:60])
        for text in expected_regular:
            if text.strip() and not any(r.text == text and r.bold is False for r in runs):
                errors.append('DOCX missing regular explanation segment: ' + text[:60])
    return errors

def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('json', type=Path); p.add_argument('--docx', type=Path)
    a = p.parse_args()
    try:
        data = json.loads(a.json.read_text(encoding='utf-8'))
        errors = validate(data, a.json.resolve().parent, a.docx)
    except (OSError, ValueError, TypeError, KeyError) as exc:
        errors = [str(exc)]
    if errors:
        for error in errors: print('ERROR:', error)
        raise SystemExit(1)
    print('Guide validation passed. Visual and language review still required.')
if __name__ == '__main__': main()
