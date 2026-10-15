#!/usr/bin/env python3
"""Build a teaching DOCX from reviewed JSON and original slide images."""
import argparse
import json
from pathlib import Path
from validate_guide import validate

def build(data, base, output, cjk_font):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.65)
    sec.left_margin = sec.right_margin = Inches(.7)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption']:
        st = doc.styles[name]
        st.font.name = 'Calibri'; st.font.color.rgb = RGBColor(0, 0, 0)
        rf = st.element.get_or_add_rPr().get_or_add_rFonts()
        rf.set(qn('w:eastAsia'), cjk_font)
        st.paragraph_format.space_after = Pt(5)
    for st in doc.styles:
        for border in list(st.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.1
    for name, size in [('Title', 26), ('Heading 1', 18), ('Heading 2', 14), ('Heading 3', 11.5)]:
        st = doc.styles[name]; st.font.size = Pt(size)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.space_before = Pt(8)
    foot = sec.footer.paragraphs[0]; foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); foot._p.append(fld)
    doc.add_paragraph(data.get('title', 'English Teaching Guide'), 'Title')
    if data.get('subtitle'): doc.add_paragraph(data['subtitle'], 'Subtitle')
    doc.add_paragraph('每页先展示原图，再给出优化正文与英文讲稿。讲稿中加粗内容对应优化后的正文，常规字体用于过渡、解释、例子和读图。')
    notes = []
    for slide in data['slides']:
        h = doc.add_paragraph(f"Slide {slide['number']:02d} — {slide['title']}", 'Heading 1')
        h.paragraph_format.page_break_before = True
        pic = doc.add_paragraph(); pic.add_run().add_picture(str(base / slide['image']), width=Inches(6.9))
        doc.add_paragraph(f"原 PPT · Slide {slide['number']:02d}", 'Caption')
        doc.add_paragraph('一、优化后的标题与正文（结构不变）', 'Heading 2')
        doc.add_paragraph('标题：' + slide['title'])
        for mod in slide['modules']:
            doc.add_paragraph(mod['label'], 'Heading 3')
            for text in mod['revised']: doc.add_paragraph(text)
        doc.add_paragraph('二、英文讲稿（按原顺序，串联图片讲解）', 'Heading 2')
        if slide.get('intro'): doc.add_paragraph(slide['intro'])
        for mod in slide['modules']:
            doc.add_paragraph(mod['label'], 'Heading 3')
            for parts in mod['script']:
                para = doc.add_paragraph()
                para.paragraph_format.widow_control = True
                for part in parts:
                    run = para.add_run(part['text']); run.bold = part['source']
        if slide.get('closing'):
            # Keep the close attached without creating a standalone heading/tail page.
            doc.paragraphs[-1].add_run(' ' + slide['closing']).bold = False
        if slide.get('instructor_notes'): notes.append((slide['number'], slide['instructor_notes']))
    if notes:
        h = doc.add_paragraph('授课校核', 'Heading 1'); h.paragraph_format.page_break_before = True
        for n, texts in notes:
            doc.add_paragraph(f'Slide {n:02d}', 'Heading 3')
            for text in texts: doc.add_paragraph(text)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)

def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('json', type=Path); p.add_argument('output', type=Path)
    p.add_argument('--cjk-font', default='Noto Sans CJK SC')
    a = p.parse_args()
    data = json.loads(a.json.read_text(encoding='utf-8')); base = a.json.resolve().parent
    errors = validate(data, base)
    if errors: raise SystemExit('\n'.join(errors))
    build(data, base, a.output, a.cjk_font)
    print(a.output.resolve())
if __name__ == '__main__': main()
