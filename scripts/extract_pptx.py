#!/usr/bin/env python3
"""Extract editable content and geometry; image text still needs visual review."""
import argparse
import json
from pathlib import Path

def extract(path):
    from pptx import Presentation
    deck = Presentation(path)
    def shapes(items):
        for shape in items:
            if shape.shape_type == 6:
                yield {'kind': 'group', 'children': list(shapes(shape.shapes))}
                continue
            box = [int(shape.left or 0), int(shape.top or 0), int(shape.width or 0), int(shape.height or 0)]
            if shape.has_text_frame and shape.text.strip():
                yield {'kind': 'text', 'box_emu': box, 'text': shape.text}
            if shape.has_table:
                yield {'kind': 'table', 'box_emu': box, 'rows': [[c.text for c in r.cells] for r in shape.table.rows]}
            if shape.shape_type == 13:
                yield {'kind': 'image', 'box_emu': box, 'note': 'Inspect rendered page for labels and diagram content.'}
    return {'source_slide_count': len(deck.slides), 'width_emu': int(deck.slide_width),
            'height_emu': int(deck.slide_height), 'slides': [
                {'number': n, 'elements': list(shapes(slide.shapes))}
                for n, slide in enumerate(deck.slides, 1)]}

def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('pptx', type=Path); p.add_argument('output', type=Path)
    a = p.parse_args()
    data = extract(a.pptx)
    a.output.parent.mkdir(parents=True, exist_ok=True)
    a.output.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f"Extracted {data['source_slide_count']} slides")
if __name__ == '__main__': main()
