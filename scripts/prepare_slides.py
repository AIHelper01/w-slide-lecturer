#!/usr/bin/env python3
"""Prepare PDF/PPTX evidence for a visual-module-aligned lecture-writing agent.

Outputs editable text + coordinates, *not* final visual modules or narration.
Rasterized diagram text requires manual / vision-capable review.
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


def sanitize_line(s: str) -> str:
    """Normalize whitespace within ONE already identified paragraph only."""
    s = s.replace("\u00a0", " ").replace("\u200b", "").replace("\u2028", "\v")
    # PPT manual line break (vertical tab) often splits one logical sentence.
    s = re.sub(r"\s*\v\s*", " ", s)
    s = re.sub(r"[ \t]+", " ", s)
    return s.strip()


def get_shape_pos(shape: Any) -> dict[str, float]:
    return {k: round((getattr(shape, k, 0) or 0) / 914400.0, 3)
            for k in ("left", "top", "width", "height")}


def read_textframe(frame: Any) -> tuple[str, str]:
    """Keep separate PPTX paragraphs intact; join only soft breaks inside a paragraph."""
    original = frame.text or ""
    ps = []
    for p in frame.paragraphs:
        val = sanitize_line(p.text or "")
        if val:
            # Treat paragraph list level as metadata, not fabricated bullet symbols.
            ps.append(val)
    return original, "\n".join(ps)


def process_shapes(shapes: Any, items: list[dict], stats: dict, prefix: str = "") -> None:
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for shape in shapes:
        shp_id = str(getattr(shape, "shape_id", "?"))
        path = f"{prefix}/{shp_id}".lstrip("/")
        st = getattr(shape, "shape_type", None)
        if st == MSO_SHAPE_TYPE.GROUP:
            stats["groups"] += 1
            process_shapes(shape.shapes, items, stats, path)
            continue
        if st == MSO_SHAPE_TYPE.PICTURE:
            stats["pictures"] += 1
        pos = get_shape_pos(shape)
        if getattr(shape, "has_text_frame", False):
            raw, clean = read_textframe(shape.text_frame)
            if clean:
                items.append({"id": path, "kind": "text", "box_inches": pos,
                              "raw": raw, "clean_paragraphs": clean})
        if getattr(shape, "has_table", False):
            stats["tables"] += 1
            for row_id, row in enumerate(shape.table.rows, start=1):
                for col_id, cell in enumerate(row.cells, start=1):
                    raw, clean = read_textframe(cell.text_frame)
                    if clean:
                        items.append({"id": f"{path}/r{row_id}c{col_id}",
                                      "kind": "table-cell", "row": row_id, "column": col_id,
                                      "box_inches": pos, "raw": raw, "clean_paragraphs": clean})
        if getattr(shape, "has_chart", False):
            stats["charts"] += 1
            try:
                chart = shape.chart
                if chart.has_title:
                    raw, clean = read_textframe(chart.chart_title.text_frame)
                    if clean:
                        items.append({"id": f"{path}/chart-title", "kind": "chart-title",
                                      "box_inches": pos, "raw": raw, "clean_paragraphs": clean})
            except Exception as exc:
                stats["warnings"].append(f"Failed to read chart in {path}: {exc}")


def notes_for_slide(slide: Any) -> str:
    try:
        if slide.has_notes_slide:
            fr = slide.notes_slide.notes_text_frame
            return fr.text if fr else ""
    except Exception:
        pass
    return ""


def manifest_pptx(src: Path) -> dict:
    from pptx import Presentation
    prs = Presentation(str(src))
    pages = []
    for number, slide in enumerate(prs.slides, 1):
        items: list[dict] = []
        stats = {"pictures": 0, "groups": 0, "tables": 0, "charts": 0, "warnings": []}
        process_shapes(slide.shapes, items, stats)
        items.sort(key=lambda item: (item["box_inches"]["top"], item["box_inches"]["left"]))
        title = ""
        try:
            if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
                title = sanitize_line(slide.shapes.title.text)
        except Exception:
            pass
        if not title:
            title = next((x["clean_paragraphs"].splitlines()[0] for x in items if x["kind"] == "text"), f"Slide {number:02d}")
        pages.append({"number": number, "title_candidate": title,
                      "text_elements_position_sorted_not_modules": items,
                      "speaker_notes_not_visible": notes_for_slide(slide),
                      "requires_visual_module_review": True, "counts": stats})
    return {"source": src.name, "format": ".pptx", "created_utc": datetime.now(timezone.utc).isoformat(),
            "slide_count": len(pages), "coordinate_system": "inches from slide top-left (groups may require visual verification)",
            "warning": "Position sorting cannot reliably group visual modules or read raster image text.", "slides": pages}


def manifest_pdf(src: Path) -> dict:
    import fitz
    doc = fitz.open(str(src))
    pages = []
    for number, page in enumerate(doc, 1):
        blocks = page.get_text("blocks")
        items = []
        for k, b in enumerate(blocks):
            if len(b) < 5 or not str(b[4]).strip():
                continue
            clean = "\n".join(sanitize_line(ln) for ln in b[4].splitlines() if sanitize_line(ln))
            items.append({"id": f"block-{k}", "kind": "pdf-text-block", "box_points": {
                "left": round(b[0], 2), "top": round(b[1], 2),
                "width": round(b[2] - b[0], 2), "height": round(b[3] - b[1], 2)},
                          "raw": b[4], "clean_paragraphs": clean})
        items.sort(key=lambda x: (x["box_points"]["top"], x["box_points"]["left"]))
        title = next((str(x["clean_paragraphs"]).splitlines()[0] for x in items), f"Slide {number:02d}")
        pages.append({"number": number, "title_candidate": title,
                      "text_elements_position_sorted_not_modules": items,
                      "speaker_notes_not_visible": "", "requires_visual_module_review": True,
                      "counts": {"pictures": len(page.get_images()), "groups": 0, "tables": 0, "charts": 0, "warnings": []}})
    doc.close()
    return {"source": src.name, "format": ".pdf", "created_utc": datetime.now(timezone.utc).isoformat(),
            "slide_count": len(pages), "coordinate_system": "PDF points from page top-left",
            "warning": "PDF reading order and line wraps must be checked visually before grouping modules.", "slides": pages}


def render_pdf(pdf: Path, out: Path, dpi: int) -> int:
    import fitz
    doc = fitz.open(str(pdf))
    imgdir = out / "slide_images"
    imgdir.mkdir(parents=True, exist_ok=True)
    for old in imgdir.glob("slide_*.png"):
        old.unlink()
    scale = dpi / 72
    for number, p in enumerate(doc, 1):
        img = p.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
        img.save(str(imgdir / f"slide_{number:02d}.png"))
    num = len(doc)
    doc.close()
    return num


def render_pptx(src: Path, out: Path, dpi: int) -> tuple[bool, str, int]:
    program = shutil.which("libreoffice") or shutil.which("soffice")
    if not program:
        return False, "LibreOffice/soffice is unavailable", 0
    try:
        import fitz  # noqa: F401
    except ImportError:
        return False, "PyMuPDF unavailable", 0
    with tempfile.TemporaryDirectory(prefix="wslide-") as tmp:
        out_tmp = Path(tmp)
        profile = (out_tmp / "lo-profile").resolve().as_uri()
        cmd = [program, f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf",
               "--outdir", str(out_tmp), str(src.resolve())]
        try:
            proc = subprocess.run(cmd, text=True, capture_output=True, timeout=240)
        except subprocess.TimeoutExpired:
            return False, "LibreOffice conversion timed out", 0
        pdf = out_tmp / f"{src.stem}.pdf"
        if proc.returncode != 0 or not pdf.exists():
            return False, (proc.stderr or proc.stdout)[-500:], 0
        return True, "Rendered through LibreOffice", render_pdf(pdf, out, dpi)


def readable_md(manifest: dict) -> str:
    lines = ["# 原始教材文字提取（供视觉核对，不是最终讲稿）", "",
             f"- 文件：{manifest['source']}", f"- 实际幻灯片：{manifest['slide_count']}",
             "- 元素按位置粗排，**尚未按真实视觉模块分组**；必须对照截图修正多栏串行、图中文字及异常换行。", ""]
    for pg in manifest["slides"]:
        lines.extend([f"## Slide {pg['number']:02d} — {pg['title_candidate']}", ""])
        for el in pg["text_elements_position_sorted_not_modules"]:
            box = el.get("box_inches", el.get("box_points"))
            loc = "in" if "box_inches" in el else "pt"
            lines.extend([f"**{el['kind']}** `id={el['id']}` `x={box['left']}{loc}, y={box['top']}{loc}`", "",
                          "```text", el["clean_paragraphs"], "```", ""])
        if pg["speaker_notes_not_visible"].strip():
            lines.extend(["**原始 PPT 备注（非画面文字）**", "", "```text",
                          pg["speaker_notes_not_visible"], "```", ""])
    return "\n".join(lines)


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("source", type=Path, help="Source .pptx or .pdf")
    p.add_argument("--out", required=True, type=Path, help="Output work folder")
    p.add_argument("--render", action="store_true", help="Render slide_images PNG files")
    p.add_argument("--dpi", type=int, default=130)
    args = p.parse_args()
    src = args.source
    if not src.is_file():
        p.error(f"Not found: {src}")
    if src.suffix.lower() not in {".pptx", ".pdf"}:
        p.error("Supports .pptx and .pdf")
    args.out.mkdir(parents=True, exist_ok=True)
    try:
        if src.suffix.lower() == ".pptx":
            manifest = manifest_pptx(src)
        else:
            manifest = manifest_pdf(src)
    except ImportError as exc:
        raise SystemExit(f"Missing dependency: {exc}. Install python-pptx and pymupdf") from exc
    (args.out / "source_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    (args.out / "selectable_text.md").write_text(readable_md(manifest), encoding="utf-8")
    print(f"Extracted {manifest['slide_count']} physical slides -> {args.out}")
    if args.render:
        if src.suffix.lower() == ".pdf":
            try:
                rendered = render_pdf(src, args.out, args.dpi)
                success, msg = True, "Rendered PDF directly"
            except ImportError as exc:
                success, msg, rendered = False, f"Missing PyMuPDF: {exc}", 0
        else:
            success, msg, rendered = render_pptx(src, args.out, args.dpi)
        print(f"Visual render: {msg} ({rendered} slides)", file=sys.stderr)
        if not success or rendered != manifest["slide_count"]:
            print("WARNING: no complete visual confirmation; do not claim to have extracted image text.", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
